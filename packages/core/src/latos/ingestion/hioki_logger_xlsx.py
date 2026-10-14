"""HIOKI Memory HiLogger recordings saved as an Excel workbook.

The logger's PC software saves a recording as a workbook as well as the text
form that `parsers.hioki_logger_csv` reads. The layout below was taken from
real exports, not from the manual::

    row 1    Trigger Time 1/2/2026 09:00:00s
    row 2    Time | CH1-A-1[°C] | CH1-A-3[°C] | CH1-A-10[V] | Event
    row 3+   1/2/2026 09:00:00.0s | 25.0 | 25.1 | 0.001234 | 0

One sheet, one row per sample. The `Event` column is absent from recordings
made without event marks; where present it is 0 on every row except the rows
where the operator pressed the mark key, which hold the running number 1, 2, …

This is a reader, not a registered parser. It hands back the recording as typed
arrays for a caller that already knows what the file is, and it stays out of
the parser registry on purpose: registered, it would compete for every workbook
a folder scan finds.

What it keeps, and what it refuses to do
----------------------------------------
* **Channel names and units are kept exactly as the header wrote them.** The
  bracketed unit is the authority, as it is for the text export: a channel on
  the 100 mV range still writes ``[V]``. Nothing is converted here.
* **No channel is given a role.** Nothing in the file says which thermocouple
  sat on the hot side, and column position is not evidence of it. A caller
  names the channels it wants; this module never guesses.
* **Every time cell is kept verbatim**, next to the elapsed seconds derived
  from it. Elapsed time comes from the time of day alone. The date part is
  compared, never parsed, so the reader does not have to decide whether the
  export wrote month/day or day/month: it is used only to tell a recording that
  crossed midnight from a clock that ran backwards.
* **Event marks are kept with the numbers the logger wrote.** Whether those
  numbers make sense as a sequence is for the caller to judge.
* **The file is opened read-only and is never written to.** Its SHA-256 is
  taken from its bytes, so a record made from it can say exactly which file it
  was.
* **A layout that does not match is refused, not repaired.** A blank cell, a
  text cell where a reading belongs, a time that does not advance, a second
  sheet: each is reported with the row it was found on. A trace with a silently
  patched time axis would move every boundary read off it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

import numpy as np
import openpyxl

from latos.core.exceptions import CorruptedFileError
from latos.ingestion.hashing import hash_file

__all__ = ["LoggerChannel", "LoggerEvent", "LoggerTrace", "read_hioki_xlsx"]

_TIME_HEADER = "time"
_EVENT_HEADER = "event"

# `CH1-A-10[V]`: the channel's name, then its unit in square brackets.
_CHANNEL_RE = re.compile(r"^(?P<name>[^\[\]]+)\[(?P<unit>[^\[\]]+)\]$")

# `1/2/2026 09:00:00.0s`: a date, then a time of day with an `s` suffix.
_TIME_RE = re.compile(r"^(?P<date>\S+)\s+(?P<h>\d{1,2}):(?P<m>\d{2}):(?P<s>\d{2}(?:\.\d+)?)s$")

_SECONDS_PER_DAY = 86400
# The header sits on the second row of every export seen. Searching a little
# further costs nothing and tolerates an extra comment row above it.
_HEADER_SEARCH_ROWS = 20
_MIN_SAMPLES = 2


@dataclass(frozen=True, slots=True)
class LoggerChannel:
    """One recorded channel, exactly as the workbook names it.

    Attributes:
        name: The channel name from the header, such as ``CH1-A-10``.
        unit: The bracketed unit from the header, such as ``V`` or ``°C``.
        header: The whole header cell.
        values: One reading per sample, in `unit`. Read-only.
    """

    name: str
    unit: str
    header: str
    values: np.ndarray


@dataclass(frozen=True, slots=True)
class LoggerEvent:
    """One event mark.

    Attributes:
        number: The number the logger wrote for the mark.
        time_s: Seconds from the first sample to the marked sample.
        row: Zero-based index of the marked sample.
    """

    number: int
    time_s: float
    row: int


@dataclass(frozen=True, slots=True)
class LoggerTrace:
    """One recording.

    Attributes:
        source: The workbook that was read.
        sha256: SHA-256 of the workbook's bytes.
        sheet: The sheet the samples came from.
        preamble: The text of every row above the header, such as the trigger
            time, verbatim.
        timestamps: The time cell of every sample, verbatim.
        time_s: Seconds since the first sample. Read-only.
        channels: Every recorded channel, in column order.
        has_event_column: Whether the recording carries an event column at all.
            False is not the same as "no marks were pressed".
        events: Every event mark, in row order.
        sample_interval_s: The median spacing between samples.
        ignored_columns: Headers of columns that are neither the time, a
            channel nor the event column. Named so nothing is dropped unseen.
    """

    source: Path
    sha256: str
    sheet: str
    preamble: tuple[str, ...]
    timestamps: tuple[str, ...]
    time_s: np.ndarray
    channels: tuple[LoggerChannel, ...]
    has_event_column: bool
    events: tuple[LoggerEvent, ...]
    sample_interval_s: float
    ignored_columns: tuple[str, ...] = ()

    def channel(self, name: str) -> LoggerChannel:
        """The channel recorded under exactly `name`.

        Raises:
            KeyError: If no channel has that name. The message lists the names
                that were recorded.
        """
        for channel in self.channels:
            if channel.name == name:
                return channel
        recorded = ", ".join(c.name for c in self.channels)
        raise KeyError(f"no channel named {name!r} in {self.source.name}; recorded: {recorded}")


def read_hioki_xlsx(path: Path, *, sheet: str | None = None) -> LoggerTrace:
    """Read one logger recording from an Excel export.

    Args:
        path: The workbook.
        sheet: The sheet to read. May be omitted when the workbook has exactly
            one sheet, which is what the logger software writes for a single
            recording.

    Returns:
        The recording, with nothing converted and no channel given a role.

    Raises:
        FileNotFoundError: If `path` does not exist.
        CorruptedFileError: If the workbook cannot be opened, has more than one
            sheet and none was named, or does not have the layout described in
            this module's docstring.
    """
    path = Path(path)
    digest = hash_file(path)
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CorruptedFileError(f"{path.name}: not a readable Excel workbook ({exc})") from exc
    try:
        sheet_name = _choose_sheet(path, list(workbook.sheetnames), sheet)
        # Read-only sheets are read lazily, so a damaged file can open cleanly
        # and fail only here.
        rows = [tuple(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    except CorruptedFileError:
        raise
    except Exception as exc:
        raise CorruptedFileError(f"{path.name}: could not read the rows ({exc})") from exc
    finally:
        workbook.close()

    header_at = _find_header(path, rows)
    header = rows[header_at]
    preamble = tuple(
        " ".join(str(cell) for cell in row if cell is not None) for row in rows[:header_at]
    )
    data = _data_rows(path, rows, header_at)

    time_column = 0
    channel_columns: list[tuple[int, str, str, str]] = []
    event_column: int | None = None
    ignored: list[str] = []
    for index, cell in enumerate(header):
        if index == time_column or cell is None:
            continue
        text = str(cell).strip()
        match = _CHANNEL_RE.match(text)
        if match:
            channel_columns.append(
                (index, match.group("name").strip(), match.group("unit").strip(), text)
            )
        elif text.lower() == _EVENT_HEADER:
            event_column = index
        elif text:
            ignored.append(text)
    if not channel_columns:
        raise CorruptedFileError(
            f"{path.name}: the header row names no channel; expected cells such as 'CH1-A-1[V]'"
        )
    names = [name for _, name, _, _ in channel_columns]
    if len(set(names)) != len(names):
        raise CorruptedFileError(f"{path.name}: the header names a channel more than once")

    timestamps, time_s = _times(path, data, time_column, header_at)
    channels = tuple(
        LoggerChannel(
            name=name,
            unit=unit,
            header=text,
            values=_readings(path, data, index, text, header_at),
        )
        for index, name, unit, text in channel_columns
    )
    events = () if event_column is None else _events(path, data, event_column, time_s, header_at)
    return LoggerTrace(
        source=path,
        sha256=digest,
        sheet=sheet_name,
        preamble=preamble,
        timestamps=timestamps,
        time_s=time_s,
        channels=channels,
        has_event_column=event_column is not None,
        events=events,
        sample_interval_s=float(np.median(np.diff(time_s))),
        ignored_columns=tuple(ignored),
    )


def _choose_sheet(path: Path, sheets: list[str], wanted: str | None) -> str:
    if wanted is not None:
        if wanted not in sheets:
            raise CorruptedFileError(
                f"{path.name}: no sheet named {wanted!r}; the workbook has {', '.join(sheets)}"
            )
        return wanted
    if len(sheets) != 1:
        # Several sheets are several recordings. Picking the first would be a
        # guess about which one was meant.
        raise CorruptedFileError(
            f"{path.name}: the workbook has {len(sheets)} sheets ({', '.join(sheets)}); "
            "name the one to read"
        )
    return sheets[0]


def _find_header(path: Path, rows: list[tuple[Any, ...]]) -> int:
    for index, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        first = row[0] if row else None
        if isinstance(first, str) and first.strip().lower() == _TIME_HEADER:
            return index
    raise CorruptedFileError(
        f"{path.name}: no header row found; expected a row whose first cell is 'Time' "
        f"within the first {_HEADER_SEARCH_ROWS} rows"
    )


def _row_number(index: int, header_at: int) -> int:
    """The spreadsheet row number of data row `index`, for messages."""
    return header_at + index + 2


def _data_rows(path: Path, rows: list[tuple[Any, ...]], header_at: int) -> list[tuple[Any, ...]]:
    data = rows[header_at + 1 :]
    # Trailing blank rows are padding; a blank row with samples after it is a
    # hole in the recording.
    while data and all(cell is None for cell in data[-1]):
        data.pop()
    for index, row in enumerate(data):
        if all(cell is None for cell in row):
            raise CorruptedFileError(
                f"{path.name}: row {_row_number(index, header_at)} is blank, with samples after it"
            )
    if len(data) < _MIN_SAMPLES:
        raise CorruptedFileError(f"{path.name}: fewer than {_MIN_SAMPLES} samples below the header")
    return data


def _times(
    path: Path, data: list[tuple[Any, ...]], column: int, header_at: int
) -> tuple[tuple[str, ...], np.ndarray]:
    """Every time cell verbatim, and the seconds elapsed since the first."""
    stamps: list[str] = []
    elapsed: list[float] = []
    first: Decimal | None = None
    previous: Decimal | None = None
    previous_date = ""
    days = 0
    for index, row in enumerate(data):
        cell = row[column] if column < len(row) else None
        match = _TIME_RE.match(cell.strip()) if isinstance(cell, str) else None
        if match is None:
            raise CorruptedFileError(
                f"{path.name}: row {_row_number(index, header_at)} has time {cell!r}; "
                "expected text such as '1/2/2026 09:00:00.0s'"
            )
        # Decimal, so 29.2 s minus 29.0 s is 0.2 s and not 0.19999999999.
        of_day = (
            Decimal(match.group("h")) * 3600
            + Decimal(match.group("m")) * 60
            + Decimal(match.group("s"))
        )
        date = match.group("date")
        if previous is not None and of_day <= previous:
            if date == previous_date:
                raise CorruptedFileError(
                    f"{path.name}: the time does not advance at row "
                    f"{_row_number(index, header_at)} ({cell!r})"
                )
            days += 1
        if first is None:
            first = of_day
        stamps.append(str(cell))
        elapsed.append(float(of_day + days * _SECONDS_PER_DAY - first))
        previous, previous_date = of_day, date
    time_s = np.asarray(elapsed, dtype=np.float64)
    time_s.setflags(write=False)
    return tuple(stamps), time_s


def _readings(
    path: Path, data: list[tuple[Any, ...]], column: int, header: str, header_at: int
) -> np.ndarray:
    values = np.empty(len(data), dtype=np.float64)
    for index, row in enumerate(data):
        cell = row[column] if column < len(row) else None
        if isinstance(cell, bool) or not isinstance(cell, (int, float)):
            raise CorruptedFileError(
                f"{path.name}: row {_row_number(index, header_at)} of {header!r} holds "
                f"{cell!r}, not a reading"
            )
        values[index] = float(cell)
    if not np.all(np.isfinite(values)):
        raise CorruptedFileError(f"{path.name}: {header!r} holds a non-finite reading")
    values.setflags(write=False)
    return values


def _events(
    path: Path,
    data: list[tuple[Any, ...]],
    column: int,
    time_s: np.ndarray,
    header_at: int,
) -> tuple[LoggerEvent, ...]:
    events: list[LoggerEvent] = []
    for index, row in enumerate(data):
        cell = row[column] if column < len(row) else None
        if cell is None or cell == 0:
            continue
        whole = (
            not isinstance(cell, bool)
            and isinstance(cell, (int, float))
            and float(cell).is_integer()
            and cell > 0
        )
        if not whole:
            raise CorruptedFileError(
                f"{path.name}: row {_row_number(index, header_at)} of the event column holds "
                f"{cell!r}, not an event number"
            )
        events.append(LoggerEvent(number=int(cell), time_s=float(time_s[index]), row=index))
    return tuple(events)
